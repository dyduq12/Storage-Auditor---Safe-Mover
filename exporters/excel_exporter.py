"""Exportador de relatórios em formato Excel (.xlsx) com openpyxl (Fase 2).

Gera uma pasta de trabalho profissional consolidando as 5 abas da auditoria:
- Resumo
- Duplicatas
- Pendentes de Backup
- Documentos Avulsos
- Jogos e Saves

Recursos implementados:
- Congelamento da primeira linha (freeze_panes = 'A2')
- Filtros automáticos em todas as colunas de dados (auto_filter)
- Cabeçalhos estilizados com preenchimento, bordas sutis e tipografia em negrito
- Formatação de tamanhos em bytes e unidades legíveis (GB/MB/KB)
- Sinalização explícita de nível de confiança (amostral vs completo)
- Sinalização 'Vínculo heurístico — revisar' para saves (destacando compartilhados e suspeitos)
"""

from __future__ import annotations

import datetime
import json
import os
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def format_bytes(byte_count: int | float | None) -> str:
    """Formata valor em bytes para string legível (B, KB, MB, GB, TB)."""
    if byte_count is None:
        return "0 B"
    bytes_val = float(byte_count)
    if bytes_val < 0:
        return "0 B"
    for unit in ["B", "KB", "MB", "GB", "TB"]:
        if bytes_val < 1024.0 or unit == "TB":
            return f"{bytes_val:.2f} {unit}" if unit != "B" else f"{int(bytes_val)} B"
        bytes_val /= 1024.0
    return f"{bytes_val:.2f} TB"


def format_timestamp(val: Any) -> str:
    """Formata timestamps ISO ou UNIX float para 'YYYY-MM-DD HH:MM:SS'."""
    if not val:
        return ""
    if isinstance(val, (int, float)):
        try:
            dt = datetime.datetime.fromtimestamp(val)
            return dt.strftime("%Y-%m-%d %H:%M:%S")
        except (ValueError, OSError, OverflowError):
            return str(val)
    if isinstance(val, str):
        try:
            # Tenta parsear formato ISO
            dt = datetime.datetime.fromisoformat(val)
            return dt.strftime("%Y-%m-%d %H:%M:%S")
        except ValueError:
            return val
    return str(val)


def classify_save_status(save_path: str, game_name: str, shared_counts: Dict[str, int]) -> str:
    """Classifica o status do vínculo do save para auditoria transparente."""
    p_lower = save_path.lower()
    g_lower = game_name.lower()
    motivos = []

    if shared_counts.get(save_path, 0) > 1:
        motivos.append("Pasta compartilhada entre múltiplos jogos/emuladores")
    if p_lower.endswith("\\steam"):
        motivos.append("Pasta raiz genérica de emulador Steam")
    if "\\fonts\\" in p_lower or "edition" in p_lower and "edition" in g_lower:
        if "\\fonts\\" in p_lower:
            motivos.append("Casamento falso por palavra 'Edition'")
    if "\\nova pasta" in p_lower or "\\light" in p_lower:
        if "light" in g_lower and not ("dying light" in p_lower):
            motivos.append("Casamento falso por termo genérico 'light'")
    if "mods" in p_lower and "expanded" in p_lower:
        motivos.append("Diretório de Mods, não savegame real")
    if "breakpoint" in p_lower and "wildlands" in g_lower:
        motivos.append("Casamento com outro jogo da franquia (Breakpoint vs Wildlands)")
    if "division" in p_lower and "wildlands" in g_lower:
        motivos.append("Casamento com outro jogo da franquia (The Division)")
    if "future soldier" in p_lower and "wildlands" in g_lower:
        motivos.append("Casamento com outro jogo da franquia (Future Soldier)")

    if motivos:
        return f"Vínculo Heurístico — Revisar ({'; '.join(motivos)})"
    return "Legítimo / Provável"


class ExcelReportExporter:
    """Exportador de dados da auditoria para planilhas Microsoft Excel (.xlsx)."""

    def __init__(self, output_dir: str = "reports") -> None:
        self.output_dir = output_dir

        # Estilos corporativos modernos
        self.header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        self.header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        self.title_font = Font(name="Segoe UI", size=14, bold=True, color="1F4E78")
        self.subtitle_font = Font(name="Segoe UI", size=10, italic=True, color="595959")
        self.data_font = Font(name="Segoe UI", size=10)
        self.alert_font = Font(name="Segoe UI", size=10, bold=True, color="C00000")
        self.warning_font = Font(name="Segoe UI", size=10, bold=True, color="B25E00")

        self.thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )
        self.align_left = Alignment(horizontal="left", vertical="center")
        self.align_right = Alignment(horizontal="right", vertical="center")
        self.align_center = Alignment(horizontal="center", vertical="center")

    def _apply_table_styling(
        self,
        ws: openpyxl.worksheet.worksheet.Worksheet,
        header_row: int = 1,
        max_col_width: int = 70,
    ) -> None:
        """Aplica estilos de cabeçalho, congelamento de painéis, filtros e auto-largura."""
        # Congela cabeçalho
        ws.freeze_panes = f"A{header_row + 1}"

        # Aplica estilos no cabeçalho
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=header_row, column=col)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = self.align_center

        # Adiciona auto-filtro na linha de cabeçalho até a última linha
        if ws.max_row >= header_row:
            end_col = get_column_letter(ws.max_column)
            ws.auto_filter.ref = f"A{header_row}:{end_col}{ws.max_row}"

        # Ajusta larguras de coluna com amostragem rápida das primeiras 50 linhas
        col_max_lens: Dict[int, int] = {}
        sample_rows = min(ws.max_row, 50)
        for row in ws.iter_rows(min_row=1, max_row=sample_rows, values_only=True):
            for col_idx, cell_value in enumerate(row, 1):
                val_str = str(cell_value or "")
                if "\n" in val_str:
                    val_str = max(val_str.split("\n"), key=len)
                col_max_lens[col_idx] = max(col_max_lens.get(col_idx, 0), len(val_str))

        for col_idx, max_len in col_max_lens.items():
            col_letter = get_column_letter(col_idx)
            adjusted_width = min(max(max_len + 3, 12), max_col_width)
            ws.column_dimensions[col_letter].width = adjusted_width

    def export(
        self,
        resumo_data: Dict[str, Any],
        duplicatas_data: Dict[str, Any],
        pendentes_data: Dict[str, Any],
        documentos_data: Dict[str, Any],
        inventario_data: Dict[str, Any],
        output_filename: str = "relatorio_auditoria_completo.xlsx",
        max_items_per_sheet: Optional[int] = None,
    ) -> str:
        """Gera o arquivo Excel completo a partir dos dados de relatório estruturados."""
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        # 1. Aba Resumo
        self._build_sheet_resumo(wb, resumo_data)

        # 2. Aba Duplicatas
        self._build_sheet_duplicatas(wb, duplicatas_data, max_items=max_items_per_sheet)

        # 3. Aba Pendentes de Backup
        self._build_sheet_pendentes(wb, pendentes_data, max_items=max_items_per_sheet)

        # 4. Aba Documentos Avulsos
        self._build_sheet_documentos(wb, documentos_data, max_items=max_items_per_sheet)

        # 5. Aba Jogos e Saves
        self._build_sheet_jogos_saves(wb, inventario_data)

        os.makedirs(self.output_dir, exist_ok=True)
        target_path = os.path.join(self.output_dir, output_filename)
        wb.save(target_path)
        return target_path

    def export_from_reports_dir(
        self,
        reports_dir: Optional[str] = None,
        output_filename: str = "relatorio_auditoria_completo.xlsx",
        max_items_per_sheet: Optional[int] = None,
    ) -> str:
        """Lê os 5 JSONs canônicos gerados em reports/ e exporta para Excel."""
        r_dir = reports_dir or self.output_dir

        def _load(filename: str) -> Dict[str, Any]:
            path = os.path.join(r_dir, filename)
            if not os.path.exists(path):
                return {}
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)

        resumo = _load("resumo_para_ia.json")
        duplicatas = _load("duplicatas_confirmadas.json")
        pendentes = _load("pendentes_de_backup_para_d.json")
        documentos = _load("documentos_avulsos.json")
        inventario = _load("inventario_jogos_e_saves.json")

        return self.export(
            resumo_data=resumo,
            duplicatas_data=duplicatas,
            pendentes_data=pendentes,
            documentos_data=documentos,
            inventario_data=inventario,
            output_filename=output_filename,
            max_items_per_sheet=max_items_per_sheet,
        )

    # -------------------------------------------------------------------------
    # Construtores de Abas
    # -------------------------------------------------------------------------

    def _build_sheet_resumo(self, wb: openpyxl.Workbook, data: Dict[str, Any]) -> None:
        """Cria e formata a aba Resumo Executivo."""
        ws = wb.create_sheet(title="Resumo")
        ws.views.sheetView[0].showGridLines = True

        ws["A1"] = "STORAGE AUDITOR & SAFE MOVER — RESUMO EXECUTIVO"
        ws["A1"].font = self.title_font
        ws["A2"] = (
            f"Relatório gerado em: {format_timestamp(data.get('timestamp_varredura'))} | "
            f"Versão Schema: {data.get('schema_version', '1.1.0')} | "
            "Modo: 100% Somente Leitura"
        )
        ws["A2"].font = self.subtitle_font

        # Tabela de Métricas Principais
        headers = ["Métrica / Categoria", "Total de Itens", "Tamanho Formatado", "Tamanho em Bytes"]
        header_row = 4
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col_idx, value=h)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = self.align_center

        totais = data.get("totais", {})
        metricas = [
            ("Arquivos Indexados no Disco C:", totais.get("arquivos_varridos_c", 0), totais.get("tamanho_total_c_bytes", 0)),
            ("Arquivos Indexados no Disco D:", totais.get("arquivos_varridos_d", 0), totais.get("tamanho_total_d_bytes", 0)),
            ("Duplicatas Confirmadas (C vs D):", totais.get("total_duplicatas", 0), totais.get("espaco_recuperavel_bytes", 0)),
            ("Espaço Recuperável Potencial em C:", totais.get("total_duplicatas", 0), totais.get("espaco_recuperavel_bytes", 0)),
            ("Pendentes de Backup para D:", totais.get("total_pendentes_backup", 0), totais.get("tamanho_pendentes_backup_bytes", 0)),
            ("Documentos Avulsos Críticos em C:", totais.get("total_documentos_avulsos", 0), totais.get("tamanho_documentos_avulsos_bytes", 0)),
            ("Jogos Instalados Mapeados:", totais.get("total_jogos_detectados", 0), totais.get("tamanho_total_jogos_bytes", 0)),
            ("Vínculos Heurísticos de Saves em C:", totais.get("total_saves_detectados", 0), totais.get("tamanho_total_saves_bytes", 0)),
        ]

        row_idx = header_row + 1
        for label, count, b_val in metricas:
            ws.cell(row=row_idx, column=1, value=label).font = self.data_font
            c2 = ws.cell(row=row_idx, column=2, value=count)
            c2.font = self.data_font
            c2.alignment = self.align_right
            c2.number_format = "#,##0"

            c3 = ws.cell(row=row_idx, column=3, value=format_bytes(b_val))
            c3.font = self.data_font
            c3.alignment = self.align_right

            c4 = ws.cell(row=row_idx, column=4, value=b_val)
            c4.font = self.data_font
            c4.alignment = self.align_right
            c4.number_format = "#,##0"

            for col in range(1, 5):
                ws.cell(row=row_idx, column=col).border = self.thin_border
            row_idx += 1

        self._apply_table_styling(ws, header_row=header_row)

    def _build_sheet_duplicatas(
        self,
        wb: openpyxl.Workbook,
        data: Dict[str, Any],
        max_items: Optional[int] = None,
    ) -> None:
        """Cria e formata a aba Duplicatas Confirmadas."""
        ws = wb.create_sheet(title="Duplicatas")
        ws.views.sheetView[0].showGridLines = True

        headers = [
            "Caminho Origem (C:)",
            "Caminho Destino (D:)",
            "Tamanho",
            "Tamanho (Bytes)",
            "Espaço Recuperável (Bytes)",
            "Nível de Confiança",
            "Hash Amostral",
            "Hash Completo",
            "Data Modificação Origem",
            "Data Modificação Destino",
        ]
        ws.append(headers)

        itens = data.get("itens", [])
        if max_items:
            itens = itens[:max_items]

        for item in itens:
            t_bytes = item.get("tamanho_bytes", 0)
            rec_bytes = item.get("espaco_recuperavel", 0)
            conf = item.get("nivel_confianca", "completo")
            row = [
                item.get("caminho_origem", ""),
                item.get("caminho_destino", ""),
                format_bytes(t_bytes),
                t_bytes,
                rec_bytes,
                conf.upper(),
                item.get("hash_amostral", ""),
                item.get("hash_completo", ""),
                format_timestamp(item.get("mtime_origem")),
                format_timestamp(item.get("mtime_destino")),
            ]
            ws.append(row)
            curr_row = ws.max_row
            ws.cell(row=curr_row, column=3).alignment = self.align_right
            c4 = ws.cell(row=curr_row, column=4)
            c4.alignment = self.align_right
            c4.number_format = "#,##0"
            c5 = ws.cell(row=curr_row, column=5)
            c5.alignment = self.align_right
            c5.number_format = "#,##0"

            if conf.lower() == "amostral":
                conf_cell = ws.cell(row=curr_row, column=6)
                conf_cell.alignment = self.align_center
                conf_cell.font = self.warning_font

        self._apply_table_styling(ws, header_row=1)

    def _build_sheet_pendentes(
        self,
        wb: openpyxl.Workbook,
        data: Dict[str, Any],
        max_items: Optional[int] = None,
    ) -> None:
        """Cria e formata a aba Pendentes de Backup."""
        ws = wb.create_sheet(title="Pendentes de Backup")
        ws.views.sheetView[0].showGridLines = True

        headers = [
            "Caminho (C:)",
            "Tamanho",
            "Tamanho (Bytes)",
            "Extensão",
            "É Imagem ISO",
            "Data Modificação",
        ]
        ws.append(headers)

        itens = data.get("itens", [])
        if max_items:
            itens = itens[:max_items]

        for item in itens:
            t_bytes = item.get("tamanho_bytes", 0)
            row = [
                item.get("caminho", ""),
                format_bytes(t_bytes),
                t_bytes,
                item.get("extensao", ""),
                "SIM" if item.get("is_iso") else "NÃO",
                format_timestamp(item.get("mtime")),
            ]
            ws.append(row)
            curr_row = ws.max_row

            ws.cell(row=curr_row, column=2).alignment = self.align_right
            c3 = ws.cell(row=curr_row, column=3)
            c3.alignment = self.align_right
            c3.number_format = "#,##0"
            ws.cell(row=curr_row, column=5).alignment = self.align_center

        self._apply_table_styling(ws, header_row=1)

    def _build_sheet_documentos(
        self,
        wb: openpyxl.Workbook,
        data: Dict[str, Any],
        max_items: Optional[int] = None,
    ) -> None:
        """Cria e formata a aba Documentos Avulsos."""
        ws = wb.create_sheet(title="Documentos Avulsos")
        ws.views.sheetView[0].showGridLines = True

        headers = [
            "Caminho (C:)",
            "Pasta Origem",
            "Tamanho",
            "Tamanho (Bytes)",
            "Extensão",
            "Prioridade",
            "Data Modificação",
        ]
        ws.append(headers)

        itens = data.get("itens", [])
        if max_items:
            itens = itens[:max_items]
        for item in itens:
            t_bytes = item.get("tamanho_bytes", 0)
            prio = item.get("prioridade", "media")
            row = [
                item.get("caminho", ""),
                item.get("pasta_origem", ""),
                format_bytes(t_bytes),
                t_bytes,
                item.get("extensao", ""),
                prio.upper(),
                format_timestamp(item.get("mtime")),
            ]
            ws.append(row)
            curr_row = ws.max_row

            ws.cell(row=curr_row, column=3).alignment = self.align_right
            c4 = ws.cell(row=curr_row, column=4)
            c4.alignment = self.align_right
            c4.number_format = "#,##0"

            prio_cell = ws.cell(row=curr_row, column=6)
            prio_cell.alignment = self.align_center
            if prio.lower() == "alta":
                prio_cell.font = self.alert_font

        self._apply_table_styling(ws, header_row=1)

    def _build_sheet_jogos_saves(self, wb: openpyxl.Workbook, data: Dict[str, Any]) -> None:
        """Cria e formata a aba Jogos e Saves."""
        ws = wb.create_sheet(title="Jogos e Saves")
        ws.views.sheetView[0].showGridLines = True

        headers = [
            "Nome do Jogo",
            "Launcher",
            "Caminho Instalação",
            "Tamanho Instalação",
            "Tamanho Instalação (Bytes)",
            "Caminho Save em C:",
            "Tamanho Save",
            "Tamanho Save (Bytes)",
            "Qtd Arquivos Save",
            "Nível de Certeza",
            "Status do Vínculo",
            "Data Modificação Save",
        ]
        ws.append(headers)

        jogos = data.get("jogos", [])

        # Identifica saves compartilhados
        from collections import Counter
        save_paths = []
        for j in jogos:
            for s in j.get("saves_em_c", []):
                save_paths.append(s.get("caminho_save", ""))
        shared_counts = Counter(save_paths)

        for j in jogos:
            j_nome = j.get("nome", "")
            j_launcher = j.get("launcher", "")
            j_path = j.get("caminho_instalacao", "")
            j_bytes = j.get("tamanho_instalacao_bytes", 0)
            saves = j.get("saves_em_c", [])

            if not saves:
                # Jogo sem saves vinculados
                row = [
                    j_nome,
                    j_launcher,
                    j_path,
                    format_bytes(j_bytes),
                    j_bytes,
                    "[Nenhum save localizado]",
                    "0 B",
                    0,
                    0,
                    "-",
                    "N/A",
                    "",
                ]
                ws.append(row)
                curr_row = ws.max_row
                ws.cell(row=curr_row, column=4).alignment = self.align_right
                c5 = ws.cell(row=curr_row, column=5)
                c5.alignment = self.align_right
                c5.number_format = "#,##0"
                for col_idx in range(1, len(headers) + 1):
                    ws.cell(row=curr_row, column=col_idx).border = self.thin_border
            else:
                for s in saves:
                    s_path = s.get("caminho_save", "")
                    s_bytes = s.get("tamanho_bytes", 0)
                    s_files = s.get("quantidade_arquivos", 0)
                    s_cert = s.get("nivel_certeza", "media")
                    status_str = classify_save_status(s_path, j_nome, shared_counts)

                    row = [
                        j_nome,
                        j_launcher,
                        j_path,
                        format_bytes(j_bytes),
                        j_bytes,
                        s_path,
                        format_bytes(s_bytes),
                        s_bytes,
                        s_files,
                        s_cert.upper(),
                        status_str,
                        format_timestamp(s.get("mtime")),
                    ]
                    ws.append(row)
                    curr_row = ws.max_row

                    ws.cell(row=curr_row, column=4).alignment = self.align_right
                    c5 = ws.cell(row=curr_row, column=5)
                    c5.alignment = self.align_right
                    c5.number_format = "#,##0"

                    ws.cell(row=curr_row, column=7).alignment = self.align_right
                    c8 = ws.cell(row=curr_row, column=8)
                    c8.alignment = self.align_right
                    c8.number_format = "#,##0"
                    c9 = ws.cell(row=curr_row, column=9)
                    c9.alignment = self.align_right
                    c9.number_format = "#,##0"

                    # Destaque de alerta se for vínculo suspeito/revisar
                    status_cell = ws.cell(row=curr_row, column=11)
                    if "Revisar" in status_str:
                        status_cell.font = self.warning_font

                    for col_idx in range(1, len(headers) + 1):
                        ws.cell(row=curr_row, column=col_idx).border = self.thin_border

        self._apply_table_styling(ws, header_row=1)
