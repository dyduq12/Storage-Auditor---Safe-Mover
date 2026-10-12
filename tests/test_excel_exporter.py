import json
import os
import threading
import tracemalloc
import zipfile
import openpyxl
import pytest

from exporters.excel_exporter import (
    ExcelReportExporter,
    ExportCancelledError,
    format_bytes,
    format_timestamp,
    classify_save_status,
)


def test_format_helpers():
    assert format_bytes(0) == "0 B"
    assert format_bytes(1024) == "1.00 KB"
    assert format_bytes(1024 * 1024 * 50) == "50.00 MB"
    assert format_bytes(1024 * 1024 * 1024 * 2) == "2.00 GB"
    assert format_bytes(None) == "0 B"

    assert format_timestamp(1670251472.0) == "2022-12-05 11:44:32"
    assert "2026-10-12" in format_timestamp("2026-10-12T00:11:56.927820")
    assert format_timestamp(None) == ""


def test_classify_save_status():
    shared = {"C:\\Users\\Public\\Documents\\Steam": 2}
    status1 = classify_save_status("C:\\Users\\Public\\Documents\\Steam", "OPERATOR", shared)
    assert "Vínculo Heurístico — Revisar" in status1
    assert "compartilhada" in status1

    status2 = classify_save_status("C:\\Users\\TestUser\\Documents\\LEGO", "LEGO MARVEL Super Heroes", {})
    assert status2 == "Legítimo / Provável"

    status3 = classify_save_status("C:\\Users\\TestUser\\Documents\\fonts\\Edition", "Dying Light 2: Reloaded Edition", {})
    assert "Revisar" in status3


def test_excel_export_synthetic_data(tmp_path):
    output_dir = str(tmp_path / "reports_test")
    exporter = ExcelReportExporter(output_dir=output_dir)

    resumo = {
        "schema_version": "1.1.0",
        "timestamp_varredura": "2026-10-12T00:00:00",
        "totais": {
            "arquivos_varridos_c": 10,
            "arquivos_varridos_d": 20,
            "total_duplicatas": 2,
            "espaco_recuperavel_bytes": 1000,
            "total_pendentes_backup": 5,
            "tamanho_pendentes_backup_bytes": 5000,
            "total_documentos_avulsos": 1,
            "tamanho_documentos_avulsos_bytes": 200,
            "total_jogos_detectados": 1,
            "tamanho_total_jogos_bytes": 10000,
            "total_saves_detectados": 1,
            "tamanho_total_saves_bytes": 300,
        },
    }

    duplicatas = {
        "itens": [
            {
                "caminho_origem": "C:\\test\\a.txt",
                "caminho_destino": "D:\\test\\a.txt",
                "tamanho_bytes": 500,
                "espaco_recuperavel": 500,
                "nivel_confianca": "completo",
                "hash_amostral": "abc",
                "hash_completo": "abc",
                "mtime_origem": 1600000000.0,
                "mtime_destino": 1600000000.0,
            },
            {
                "caminho_origem": "C:\\test\\big.iso",
                "caminho_destino": "D:\\test\\big.iso",
                "tamanho_bytes": 500000000,
                "espaco_recuperavel": 500000000,
                "nivel_confianca": "amostral",
                "hash_amostral": "def",
                "hash_completo": None,
                "mtime_origem": 1600000000.0,
                "mtime_destino": 1600000000.0,
            },
        ]
    }

    pendentes = {
        "itens": [
            {
                "caminho": "C:\\test\\b.txt",
                "tamanho_bytes": 1000,
                "mtime": 1600000000.0,
                "extensao": ".txt",
                "is_iso": False,
            }
        ]
    }

    documentos = {
        "itens": [
            {
                "caminho": "C:\\test\\relatorio.pdf",
                "pasta_origem": "Documents",
                "tamanho_bytes": 200,
                "extensao": ".pdf",
                "prioridade": "alta",
                "mtime": 1600000000.0,
            }
        ]
    }

    inventario = {
        "jogos": [
            {
                "nome": "Super Game",
                "launcher": "Steam",
                "caminho_instalacao": "D:\\SteamLibrary\\SuperGame",
                "tamanho_instalacao_bytes": 10000,
                "saves_em_c": [
                    {
                        "caminho_save": "C:\\Users\\User\\Documents\\SuperGame",
                        "tamanho_bytes": 300,
                        "quantidade_arquivos": 3,
                        "nivel_certeza": "alta",
                        "mtime": 1600000000.0,
                    }
                ],
            }
        ]
    }

    excel_path = exporter.export(
        resumo_data=resumo,
        duplicatas_data=duplicatas,
        pendentes_data=pendentes,
        documentos_data=documentos,
        inventario_data=inventario,
        output_filename="test_output.xlsx",
    )

    assert os.path.exists(excel_path)

    # Inspeciona o arquivo gerado
    wb = openpyxl.load_workbook(excel_path)
    sheet_names = wb.sheetnames
    expected_sheets = ["Resumo", "Duplicatas", "Pendentes de Backup", "Documentos Avulsos", "Jogos e Saves"]
    for s in expected_sheets:
        assert s in sheet_names, f"Aba '{s}' não encontrada no arquivo Excel"

    # Valida aba Resumo
    ws_resumo = wb["Resumo"]
    assert ws_resumo["A1"].value == "STORAGE AUDITOR & SAFE MOVER — RESUMO EXECUTIVO"
    assert ws_resumo.auto_filter.ref is not None

    # Valida aba Duplicatas
    ws_dup = wb["Duplicatas"]
    assert ws_dup.freeze_panes == "A2"
    assert ws_dup.auto_filter.ref is not None
    assert ws_dup.cell(row=2, column=1).value == "C:\\test\\a.txt"
    assert "AMOSTRAL" in ws_dup.cell(row=3, column=6).value

    # Valida aba Pendentes
    ws_pen = wb["Pendentes de Backup"]
    assert ws_pen.freeze_panes == "A2"
    assert ws_pen.auto_filter.ref is not None

    # Valida aba Jogos e Saves
    ws_jogos = wb["Jogos e Saves"]
    assert ws_jogos.freeze_panes == "A2"
    assert ws_jogos.cell(row=2, column=1).value == "Super Game"
    assert ws_jogos.cell(row=2, column=11).value == "Legítimo / Provável"


def test_excel_export_from_real_reports(tmp_path):
    """Testa exportação lendo a pasta reports/ existente do repositório."""
    if not os.path.exists("reports/resumo_para_ia.json"):
        pytest.skip("reports/resumo_para_ia.json não existe no ambiente")

    output_dir = str(tmp_path / "real_excel_test")
    exporter = ExcelReportExporter(output_dir=output_dir)
    excel_path = exporter.export_from_reports_dir(
        reports_dir="reports",
        output_filename="auditoria_real.xlsx",
        max_items_per_sheet=200,
    )

    assert os.path.exists(excel_path)
    assert os.path.getsize(excel_path) > 1000

    wb = openpyxl.load_workbook(excel_path, read_only=True)
    expected_sheets = ["Resumo", "Duplicatas", "Pendentes de Backup", "Documentos Avulsos", "Jogos e Saves"]
    for s in expected_sheets:
        assert s in wb.sheetnames


def test_large_synthetic_export_memory_peak(tmp_path):
    """Gera massa sintética grande (>60.000 itens) e mede pico de memória com tracemalloc."""
    reports_dir = tmp_path / "large_reports"
    reports_dir.mkdir()

    # 1. Resumo
    resumo = {
        "schema_version": "1.1.0",
        "timestamp_varredura": "2026-10-12T00:00:00",
        "totais": {
            "arquivos_varridos_c": 62000,
            "arquivos_varridos_d": 62000,
            "total_duplicatas": 20000,
            "espaco_recuperavel_bytes": 20000 * 1024,
            "total_pendentes_backup": 40000,
            "tamanho_pendentes_backup_bytes": 40000 * 2048,
            "total_documentos_avulsos": 2000,
            "tamanho_documentos_avulsos_bytes": 2000 * 512,
            "total_jogos_detectados": 10,
            "tamanho_total_jogos_bytes": 10 * 1024 * 1024,
            "total_saves_detectados": 20,
            "tamanho_total_saves_bytes": 20 * 4096,
        },
    }
    with open(reports_dir / "resumo_para_ia.json", "w", encoding="utf-8") as f:
        json.dump(resumo, f)

    # 2. Duplicatas (20.000)
    with open(reports_dir / "relatorio_duplicatas.json", "w", encoding="utf-8") as f:
        f.write('{"itens": [\n')
        for i in range(20000):
            item = {
                "caminho_origem": f"C:\\synth\\orig_{i}.txt",
                "caminho_destino": f"D:\\synth\\dest_{i}.txt",
                "tamanho_bytes": 1024,
                "espaco_recuperavel": 1024,
                "nivel_confianca": "amostral" if i % 1000 == 0 else "completo",
                "hash_amostral": "hash_sample",
                "hash_completo": "hash_full" if i % 1000 != 0 else None,
                "mtime_origem": 1600000000.0,
                "mtime_destino": 1600000000.0,
            }
            f.write(json.dumps(item))
            if i < 19999:
                f.write(",\n")
        f.write("\n]}")

    # 3. Pendentes (40.000)
    with open(reports_dir / "relatorio_pendentes_backup.json", "w", encoding="utf-8") as f:
        f.write('{"itens": [\n')
        for i in range(40000):
            item = {
                "caminho": f"C:\\synth\\pend_{i}.dat",
                "tamanho_bytes": 2048,
                "mtime": 1600000000.0,
                "extensao": ".dat",
                "is_iso": False,
            }
            f.write(json.dumps(item))
            if i < 39999:
                f.write(",\n")
        f.write("\n]}")

    # 4. Documentos (2.000)
    with open(reports_dir / "relatorio_documentos_avulsos.json", "w", encoding="utf-8") as f:
        f.write('{"itens": [\n')
        for i in range(20000 if False else 2000):
            item = {
                "caminho": f"C:\\synth\\doc_{i}.pdf",
                "pasta_origem": "Documents",
                "tamanho_bytes": 512,
                "extensao": ".pdf",
                "prioridade": "alta",
                "mtime": 1600000000.0,
            }
            f.write(json.dumps(item))
            if i < 1999:
                f.write(",\n")
        f.write("\n]}")

    # 5. Jogos e Saves (10 jogos, 20 saves)
    jogos = []
    for g in range(10):
        jogos.append({
            "nome": f"Synthetic Game {g}",
            "launcher": "Steam",
            "caminho_instalacao": f"D:\\Games\\Game_{g}",
            "tamanho_instalacao_bytes": 1024 * 1024,
            "saves_em_c": [
                {
                    "caminho_save": f"C:\\Users\\User\\Documents\\Game_{g}_save_1",
                    "tamanho_bytes": 4096,
                    "quantidade_arquivos": 2,
                    "nivel_certeza": "alta",
                    "mtime": 1600000000.0,
                },
                {
                    "caminho_save": f"C:\\Users\\Public\\Documents\\Game_{g}_save_2",
                    "tamanho_bytes": 4096,
                    "quantidade_arquivos": 1,
                    "nivel_certeza": "media",
                    "mtime": 1600000000.0,
                },
            ],
        })
    with open(reports_dir / "inventario_jogos_saves.json", "w", encoding="utf-8") as f:
        json.dump({"jogos": jogos}, f)

    # Mede pico de memória com tracemalloc
    out_dir = tmp_path / "large_out"
    exporter = ExcelReportExporter(output_dir=str(out_dir))

    tracemalloc.start()
    excel_path = exporter.export_from_reports_dir(
        reports_dir=str(reports_dir),
        output_filename="relatorio_sintetico_grande.xlsx",
    )
    current_mem, peak_mem = tracemalloc.get_traced_memory()
    tracemalloc.stop()

    peak_mb = peak_mem / (1024 * 1024)
    print(f"\n[TRACE] Pico de memória medido no teste sintético de ~62.000 itens: {peak_mb:.2f} MB")

    # Garante que o pico de memória em modo write-only seja estritamente baixo (< 120 MB)
    assert peak_mb < 120.0, f"Pico de memória excessivo: {peak_mb:.2f} MB"

    # Valida integridade do arquivo gerado
    assert os.path.exists(excel_path)
    assert zipfile.is_zipfile(excel_path)

    wb = openpyxl.load_workbook(excel_path, read_only=True)
    assert len(wb.sheetnames) == 5
    for sheet in ["Resumo", "Duplicatas", "Pendentes de Backup", "Documentos Avulsos", "Jogos e Saves"]:
        assert sheet in wb.sheetnames


def test_export_cancellation_cleans_temp_file(tmp_path):
    """Testa se o cancelamento remove o arquivo temporário e preserva o XLSX anterior e os JSONs."""
    reports_dir = tmp_path / "cancel_reports"
    reports_dir.mkdir()

    # Prepara JSONs mínimos
    with open(reports_dir / "resumo_para_ia.json", "w", encoding="utf-8") as f:
        json.dump({"schema_version": "1.1.0", "totais": {}}, f)
    with open(reports_dir / "relatorio_duplicatas.json", "w", encoding="utf-8") as f:
        f.write('{"itens": [' + ','.join([f'{{"caminho_origem": "C:\\\\f{i}", "tamanho_bytes": 10}}' for i in range(1000)]) + ']}')
    with open(reports_dir / "relatorio_pendentes_backup.json", "w", encoding="utf-8") as f:
        json.dump({"itens": []}, f)
    with open(reports_dir / "relatorio_documentos_avulsos.json", "w", encoding="utf-8") as f:
        json.dump({"itens": []}, f)
    with open(reports_dir / "inventario_jogos_saves.json", "w", encoding="utf-8") as f:
        json.dump({"jogos": []}, f)

    out_dir = tmp_path / "cancel_out"
    out_dir.mkdir()
    target_xlsx = out_dir / "relatorio.xlsx"
    original_content = b"CONTEUDO_ANTERIOR_PRESERVADO"
    target_xlsx.write_bytes(original_content)

    cancel_event = threading.Event()

    def _progress_cb(msg, frac):
        # Dispara cancelamento assim que começar a emitir progresso
        cancel_event.set()

    exporter = ExcelReportExporter(output_dir=str(out_dir))
    with pytest.raises(ExportCancelledError):
        exporter.export_from_reports_dir(
            reports_dir=str(reports_dir),
            target_filepath=str(target_xlsx),
            cancel_event=cancel_event,
            progress_callback=_progress_cb,
        )

    # 1. Verifica que o arquivo existente original está intocado
    assert target_xlsx.exists()
    assert target_xlsx.read_bytes() == original_content

    # 2. Verifica que nenhum temporário ficou órfão no diretório de saída
    temp_files = [f for f in os.listdir(str(out_dir)) if f.startswith(".~tmp_")]
    assert len(temp_files) == 0, f"Arquivos temporários órfãos encontrados: {temp_files}"

    # 3. Verifica que os arquivos JSON originais continuam intactos
    assert (reports_dir / "resumo_para_ia.json").exists()
    assert (reports_dir / "relatorio_duplicatas.json").exists()


def test_export_error_cleans_temp_file(tmp_path):
    """Testa se uma exceção durante a exportação remove o arquivo temporário."""
    reports_dir = tmp_path / "err_reports"
    reports_dir.mkdir()

    with open(reports_dir / "resumo_para_ia.json", "w", encoding="utf-8") as f:
        json.dump({"schema_version": "1.1.0", "totais": {}}, f)
    # JSON de duplicatas malformado para disparar erro durante a leitura
    with open(reports_dir / "relatorio_duplicatas.json", "w", encoding="utf-8") as f:
        f.write('{"itens": [ INVALID_JSON')

    out_dir = tmp_path / "err_out"
    out_dir.mkdir()
    target_xlsx = out_dir / "relatorio.xlsx"

    exporter = ExcelReportExporter(output_dir=str(out_dir))
    with pytest.raises(Exception):
        exporter.export_from_reports_dir(
            reports_dir=str(reports_dir),
            target_filepath=str(target_xlsx),
        )

    # Verifica que nenhum temporário sobrou
    temp_files = [f for f in os.listdir(str(out_dir)) if f.startswith(".~tmp_")]
    assert len(temp_files) == 0, f"Arquivos temporários órfãos após erro: {temp_files}"
    # Target nunca foi criado
    assert not target_xlsx.exists()


def test_export_atomic_replacement_on_existing_file(tmp_path):
    """Testa substituição atômica de arquivo existente quando a exportação tem sucesso."""
    reports_dir = tmp_path / "atomic_reports"
    reports_dir.mkdir()

    with open(reports_dir / "resumo_para_ia.json", "w", encoding="utf-8") as f:
        json.dump({"schema_version": "1.1.0", "totais": {"arquivos_varridos_c": 1}}, f)
    with open(reports_dir / "relatorio_duplicatas.json", "w", encoding="utf-8") as f:
        json.dump({"itens": []}, f)
    with open(reports_dir / "relatorio_pendentes_backup.json", "w", encoding="utf-8") as f:
        json.dump({"itens": []}, f)
    with open(reports_dir / "relatorio_documentos_avulsos.json", "w", encoding="utf-8") as f:
        json.dump({"itens": []}, f)
    with open(reports_dir / "inventario_jogos_saves.json", "w", encoding="utf-8") as f:
        json.dump({"jogos": []}, f)

    out_dir = tmp_path / "atomic_out"
    out_dir.mkdir()
    target_xlsx = out_dir / "relatorio.xlsx"
    target_xlsx.write_bytes(b"DADOS_ANTIGOS")

    exporter = ExcelReportExporter(output_dir=str(out_dir))
    final_path = exporter.export_from_reports_dir(
        reports_dir=str(reports_dir),
        target_filepath=str(target_xlsx),
    )

    assert os.path.abspath(final_path) == os.path.abspath(str(target_xlsx))
    assert target_xlsx.exists()
    assert zipfile.is_zipfile(target_xlsx)
    assert target_xlsx.read_bytes() != b"DADOS_ANTIGOS"

    # Sem temporários órfãos
    temp_files = [f for f in os.listdir(str(out_dir)) if f.startswith(".~tmp_")]
    assert len(temp_files) == 0

